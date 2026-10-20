"""Tests del módulo de Reportes y Exportación (Fase 9).

Cubre: los 9 tipos de reporte (ventas, productos vendidos, inventario,
stock bajo, compras, caja, clientes con crédito, cuentas por cobrar y
utilidades), filtros y fechas, totales sobre el conjunto completo,
paginación, RBAC (reports.read / reports.export), reportes vacíos,
validaciones de parámetros y exportación real a PDF/Excel con nombres
de archivo descriptivos.
"""

import io
import re
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.reports.excel import currency_number_format
from app.reports.pdf import format_money
from tests.conftest import auth_headers

BASE = "/api/reports"

REPORT_KEYS = [
    "sales",
    "products-sold",
    "inventory",
    "low-stock",
    "purchases",
    "cash-registers",
    "customers-credit",
    "receivables",
    "utilities",
]


# ----------------------------------------------------------------------
# Helpers de datos (fabricación vía API, convención del proyecto)
# ----------------------------------------------------------------------
def _sale(
    client: TestClient,
    user,
    product,
    quantity: int = 2,
    payment_method: str = "CASH",
    customer_id=None,
) -> dict:
    payload: dict = {
        "payment_method": payment_method,
        "items": [{"product_id": str(product.id), "quantity": quantity}],
    }
    if customer_id is not None:
        payload["customer_id"] = str(customer_id)
    created = client.post("/api/sales", json=payload, headers=auth_headers(user))
    assert created.status_code == 201, created.text
    confirmed = client.post(
        f"/api/sales/{created.json()['id']}/confirm", headers=auth_headers(user)
    )
    assert confirmed.status_code == 200, confirmed.text
    return confirmed.json()


def _draft_sale(client: TestClient, user, product) -> dict:
    created = client.post(
        "/api/sales",
        json={
            "payment_method": "CASH",
            "items": [{"product_id": str(product.id), "quantity": 1}],
        },
        headers=auth_headers(user),
    )
    assert created.status_code == 201, created.text
    return created.json()


def _purchase_confirmed(
    client: TestClient, user, supplier, product
) -> dict:
    purchase = client.post(
        "/api/purchases",
        json={
            "supplier_id": str(supplier.id),
            "items": [
                {
                    "product_id": str(product.id),
                    "quantity": 10,
                    "unit_cost": "10.00",
                }
            ],
        },
        headers=auth_headers(user),
    )
    assert purchase.status_code == 201, purchase.text
    confirmed = client.post(
        f"/api/purchases/{purchase.json()['id']}/confirm",
        headers=auth_headers(user),
    )
    assert confirmed.status_code == 200, confirmed.text
    return confirmed.json()


def _purchase_draft(
    client: TestClient, user, supplier, product
) -> dict:
    purchase = client.post(
        "/api/purchases",
        json={
            "supplier_id": str(supplier.id),
            "items": [
                {
                    "product_id": str(product.id),
                    "quantity": 5,
                    "unit_cost": "7.00",
                }
            ],
        },
        headers=auth_headers(user),
    )
    assert purchase.status_code == 201, purchase.text
    return purchase.json()


def _credit_customer_with_receivable(
    client: TestClient, admin, make_customer, amount="300.00"
):
    customer = make_customer(credit_enabled=True, credit_limit="100000.00")
    receivable = client.post(
        "/api/receivables",
        json={
            "customer_id": str(customer.id),
            "original_amount": amount,
            "description": "Cuenta de prueba reportes",
        },
        headers=auth_headers(admin),
    )
    assert receivable.status_code == 201, receivable.text
    return customer, receivable.json()


@pytest.fixture(scope="function")
def admin_headers(admin_user):
    return auth_headers(admin_user)


@pytest.fixture(scope="function")
def cashier_headers(cashier_user):
    return auth_headers(cashier_user)


@pytest.fixture(scope="function")
def invmgr_headers(inventory_manager_user):
    return auth_headers(inventory_manager_user)


def _get(client, path, headers, expected=200):
    response = client.get(path, headers=headers)
    assert response.status_code == expected, response.text
    return response


# ----------------------------------------------------------------------
# RBAC
# ----------------------------------------------------------------------
def test_reports_require_authentication(client):
    response = client.get(f"{BASE}/sales")
    assert response.status_code == 401


def test_reports_json_denied_for_cashier(client, cashier_headers):
    response = client.get(f"{BASE}/sales", headers=cashier_headers)
    assert response.status_code == 403
    assert response.json()["error_code"] == "PERMISSION_DENIED"


def test_reports_export_denied_for_cashier(client, cashier_headers):
    response = client.get(
        f"{BASE}/sales/export?format=pdf", headers=cashier_headers
    )
    assert response.status_code == 403
    assert response.json()["error_code"] == "PERMISSION_DENIED"


def test_reports_json_allowed_for_admin_and_inventory_manager(
    client, admin_headers, invmgr_headers
):
    for headers in (admin_headers, invmgr_headers):
        _get(client, f"{BASE}/inventory", headers)


def test_reports_export_allowed_for_admin_and_inventory_manager(
    client, admin_headers, invmgr_headers
):
    for headers in (admin_headers, invmgr_headers):
        response = _get(
            client, f"{BASE}/inventory/export?format=pdf", headers
        )
        assert response.content.startswith(b"%PDF")


def test_reports_unknown_type_returns_404(client, admin_headers):
    response = client.get(f"{BASE}/no-existe", headers=admin_headers)
    assert response.status_code == 404
    assert response.json()["error_code"] == "REPORT_NOT_FOUND"


def test_reports_unknown_type_export_404(client, admin_headers):
    response = client.get(
        f"{BASE}/no-existe/export?format=pdf", headers=admin_headers
    )
    assert response.status_code == 404


# ----------------------------------------------------------------------
# Validaciones de parámetros
# ----------------------------------------------------------------------
def test_sales_report_rejects_filter_not_applicable(client, admin_headers):
    valid_uuid = "00000000-0000-0000-0000-000000000000"
    response = client.get(
        f"{BASE}/sales", params={"category_id": valid_uuid}, headers=admin_headers
    )
    assert response.status_code == 422
    body = response.json()
    assert body["error_code"] == "INVALID_FILTER_FOR_REPORT"
    assert "category_id" in body["detail"]


def test_report_invalid_payment_method_422(client, admin_headers):
    response = client.get(
        f"{BASE}/sales",
        params={"payment_method": "BITCOIN"},
        headers=admin_headers,
    )
    assert response.status_code == 422
    assert response.json()["error_code"] == "INVALID_FILTER_VALUE"


def test_purchases_report_invalid_status_422(client, admin_headers):
    response = client.get(
        f"{BASE}/purchases", params={"status": "WEIRD"}, headers=admin_headers
    )
    assert response.status_code == 422
    assert response.json()["error_code"] == "INVALID_FILTER_VALUE"


def test_report_date_from_after_date_to_422(client, admin_headers):
    response = client.get(
        f"{BASE}/sales",
        params={
            "date_from": "2026-08-24T00:00:00",
            "date_to": "2026-08-01T00:00:00",
        },
        headers=admin_headers,
    )
    assert response.status_code == 422
    assert response.json()["error_code"] == "INVALID_DATE_RANGE"


def test_utilities_invalid_period_422(client, admin_headers):
    response = client.get(
        f"{BASE}/utilities", params={"period": "WEEK"}, headers=admin_headers
    )
    assert response.status_code == 422
    assert response.json()["error_code"] == "INVALID_FILTER_VALUE"


def test_export_invalid_format_422(client, admin_headers):
    response = client.get(
        f"{BASE}/sales/export", params={"format": "csv"}, headers=admin_headers
    )
    assert response.status_code == 422
    assert response.json()["error_code"] == "INVALID_EXPORT_FORMAT"


# ----------------------------------------------------------------------
# Reporte de ventas
# ----------------------------------------------------------------------
def test_sales_report_includes_only_confirmed_sales(
    client, admin_user, admin_headers, make_product
):
    product = make_product()
    sale = _sale(client, admin_user, product, quantity=2)
    _draft_sale(client, admin_user, product)

    data = _get(client, f"{BASE}/sales", admin_headers).json()
    assert data["total"] == 1
    row = data["items"][0]
    assert row["sale_number"] == sale["sale_number"]
    assert row["seller"] == admin_user.username
    assert row["method_label"] == "Efectivo"
    assert row["total"] == sale["total"]
    assert data["totals"]["sales_count"] == 1
    assert data["totals"]["total_sum"] == sale["total"]
    assert data["totals"]["subtotal_sum"] == sale["subtotal"]
    labels = [c["key"] for c in data["columns"]]
    assert labels[0] == "sale_number"


def test_sales_report_filter_payment_method(
    client, admin_user, cashier_user, admin_headers, make_product
):
    product = make_product()
    cash_sale = _sale(client, admin_user, product, payment_method="CASH")
    card_sale = _sale(client, cashier_user, product, payment_method="CARD")

    all_data = client.get(f"{BASE}/sales", headers=admin_headers).json()
    assert all_data["total"] == 2

    filtered = client.get(
        f"{BASE}/sales",
        params={"payment_method": "CARD"},
        headers=admin_headers,
    )
    assert filtered.status_code == 200
    body = filtered.json()
    assert body["total"] == 1
    assert body["items"][0]["sale_number"] == card_sale["sale_number"]
    assert body["totals"]["total_sum"] == card_sale["total"]
    assert body["items"][0]["method_label"] == "Tarjeta"
    assert cash_sale["total"] != card_sale["total"] or cash_sale["id"] != card_sale["id"]


def test_sales_report_filters_exclude_by_dates(
    client, admin_user, admin_headers, make_product
):
    product = make_product()
    _sale(client, admin_user, product)

    future_from = (datetime.now() + timedelta(days=1)).isoformat()
    empty = client.get(
        f"{BASE}/sales", params={"date_from": future_from}, headers=admin_headers
    )
    assert empty.status_code == 200
    assert empty.json()["total"] == 0
    assert empty.json()["totals"]["total_sum"] == "0.00"

    yesterday = (datetime.now() - timedelta(days=1)).isoformat()
    tomorrow = (datetime.now() + timedelta(days=1)).isoformat()
    full = client.get(
        f"{BASE}/sales",
        params={"date_from": yesterday, "date_to": tomorrow},
        headers=admin_headers,
    )
    assert full.status_code == 200
    assert full.json()["total"] == 1
    applied = {f["label"]: f["value"] for f in full.json()["filters_applied"]}
    assert "Desde" in applied and "Hasta" in applied


def test_sales_report_filter_seller(
    client, admin_user, cashier_user, admin_headers, make_product
):
    product = make_product()
    _sale(client, admin_user, product, quantity=1)
    _sale(client, cashier_user, product, quantity=1)

    data = client.get(
        f"{BASE}/sales", params={"user_id": str(admin_user.id)}, headers=admin_headers
    ).json()
    assert data["total"] == 1
    assert data["items"][0]["seller"] == admin_user.username


def test_sales_report_pagination_keeps_total(
    client, admin_user, admin_headers, make_product
):
    product = make_product(stock=100)
    for _ in range(3):
        _sale(client, admin_user, product, quantity=1)

    page_one = client.get(
        f"{BASE}/sales", params={"limit": 2}, headers=admin_headers
    ).json()
    assert len(page_one["items"]) == 2
    assert page_one["total"] == 3

    page_two = client.get(
        f"{BASE}/sales",
        params={"limit": 2, "skip": 2},
        headers=admin_headers,
    ).json()
    assert len(page_two["items"]) == 1
    assert page_two["total"] == 3


def test_sales_report_empty_has_zeroed_totals(client, admin_headers):
    data = _get(client, f"{BASE}/sales", admin_headers).json()
    assert data["items"] == []
    assert data["total"] == 0
    assert data["totals"]["sales_count"] == 0
    assert data["totals"]["total_sum"] == "0.00"


# ----------------------------------------------------------------------
# Productos vendidos
# ----------------------------------------------------------------------
def test_products_sold_aggregates_quantity_and_revenue(
    client, admin_user, admin_headers, make_category, make_product
):
    category = make_category()
    cheap = make_product(category=category, sale_price="100.00", stock=50)
    pricey = make_product(category=category, sale_price="500.00", stock=50)
    _sale(client, admin_user, cheap, quantity=3)
    _sale(client, admin_user, cheap, quantity=1)
    _sale(client, admin_user, pricey, quantity=2)

    data = _get(
        client, f"{BASE}/products-sold", admin_headers
    ).json()
    assert data["total"] == 2
    first, second = data["items"]
    assert first["product_name"] == cheap.name
    assert first["quantity_sold"] == 4
    assert first["sales_count"] == 2
    assert first["revenue"] == "400.00"
    assert second["quantity_sold"] == 2
    totals = data["totals"]
    assert totals["units_sold"] == 6
    assert totals["revenue_sum"] == "1400.00"


def test_products_sold_filter_category(
    client, admin_user, admin_headers, make_category, make_product
):
    cat_a = make_category()
    cat_b = make_category()
    prod_a = make_product(category=cat_a, stock=20)
    prod_b = make_product(category=cat_b, stock=20)
    _sale(client, admin_user, prod_a)
    _sale(client, admin_user, prod_b)

    data = client.get(
        f"{BASE}/products-sold",
        params={"category_id": str(cat_a.id)},
        headers=admin_headers,
    ).json()
    assert data["total"] == 1
    assert data["items"][0]["product_name"] == prod_a.name
    assert data["items"][0]["category_name"] == cat_a.name


def test_products_sold_ignores_draft_sales(
    client, admin_user, admin_headers, make_product
):
    product = make_product(stock=20)
    _draft_sale(client, admin_user, product)
    data = _get(client, f"{BASE}/products-sold", admin_headers).json()
    assert data["total"] == 0
    assert data["totals"]["units_sold"] == 0


# ----------------------------------------------------------------------
# Inventario
# ----------------------------------------------------------------------
def test_inventory_report_totals_value_cost_and_sale(
    client, admin_headers, make_product
):
    make_product(stock=10, purchase_price="50.00", sale_price="100.00")
    make_product(stock=3, purchase_price="10.00", sale_price="25.00")

    data = _get(client, f"{BASE}/inventory", admin_headers).json()
    assert data["total"] == 2
    assert data["totals"]["products_count"] == 2
    assert data["totals"]["total_units"] == 13
    assert data["totals"]["total_value_cost"] == "530.00"
    assert data["totals"]["total_value_sale"] == "1075.00"
    row_values = {(r["stock"], r["value_cost"]) for r in data["items"]}
    assert (10, "500.00") in row_values


def test_inventory_report_filters_state_and_category(
    client, admin_headers, make_category, make_product, db_session
):
    category = make_category()
    active = make_product(name="Activo A", category=category, stock=5)
    inactive = make_product(name="Inactivo B", stock=5, is_active=False)

    only_active = client.get(
        f"{BASE}/inventory", params={"state": "active"}, headers=admin_headers
    ).json()
    assert only_active["total"] == 1
    assert only_active["items"][0]["product_name"] == active.name

    only_inactive = client.get(
        f"{BASE}/inventory", params={"state": "inactive"}, headers=admin_headers
    ).json()
    assert only_inactive["total"] == 1
    assert only_inactive["items"][0]["product_name"] == inactive.name

    by_category = client.get(
        f"{BASE}/inventory",
        params={"category_id": str(category.id)},
        headers=admin_headers,
    ).json()
    assert by_category["total"] == 1
    assert by_category["items"][0]["product_name"] == active.name


def test_low_stock_report_lists_and_orders_by_shortage(
    client, admin_headers, make_product
):
    worst = make_product(name="Crítico", stock=1, minimum_stock=5)
    mild = make_product(name="Leve", stock=3, minimum_stock=4)
    make_product(name="Sano", stock=50, minimum_stock=2)

    data = _get(client, f"{BASE}/low-stock", admin_headers).json()
    names = [r["product_name"] for r in data["items"]]
    assert names == [worst.name, mild.name]
    assert data["items"][0]["shortage"] == 4
    assert data["totals"]["products_count"] == 2
    assert data["totals"]["total_shortage"] == 5


def test_low_stock_filter_category(
    client, admin_headers, make_category, make_product
):
    cat = make_category()
    low_in_cat = make_product(category=cat, stock=0, minimum_stock=3)
    make_product(stock=0, minimum_stock=3)

    data = client.get(
        f"{BASE}/low-stock",
        params={"category_id": str(cat.id)},
        headers=admin_headers,
    ).json()
    assert [r["product_name"] for r in data["items"]] == [low_in_cat.name]
    assert data["totals"]["total_shortage"] == 3


# ----------------------------------------------------------------------
# Compras
# ----------------------------------------------------------------------
def test_purchases_report_rows_totals_and_status_filter(
    client, admin_user, admin_headers, make_supplier, make_product
):
    supplier = make_supplier()
    product = make_product(stock=0)
    confirmed = _purchase_confirmed(client, admin_user, supplier, product)
    draft = _purchase_draft(client, admin_user, supplier, product)

    all_data = _get(client, f"{BASE}/purchases", admin_headers).json()
    assert all_data["total"] == 2
    numbers = {r["purchase_number"] for r in all_data["items"]}
    assert confirmed["purchase_number"] in numbers
    assert draft["purchase_number"] in numbers

    confirmed_only = client.get(
        f"{BASE}/purchases",
        params={"status": "CONFIRMED"},
        headers=admin_headers,
    ).json()
    assert confirmed_only["total"] == 1
    assert confirmed_only["items"][0]["status_label"] == "Confirmada"
    assert confirmed_only["totals"]["total_sum"] == confirmed["total"]

    assert all_data["totals"]["purchases_count"] == 2


def test_purchases_report_filter_supplier_and_dates(
    client, admin_user, admin_headers, make_supplier, make_product
):
    supplier_a = make_supplier()
    supplier_b = make_supplier()
    product = make_product(stock=0)
    _purchase_confirmed(client, admin_user, supplier_a, product)
    _purchase_confirmed(client, admin_user, supplier_b, product)

    data = client.get(
        f"{BASE}/purchases",
        params={"supplier_id": str(supplier_a.id)},
        headers=admin_headers,
    ).json()
    assert data["total"] == 1
    assert data["items"][0]["supplier_name"] == supplier_a.name

    future = (datetime.now() + timedelta(days=1)).isoformat()
    empty = client.get(
        f"{BASE}/purchases", params={"date_from": future}, headers=admin_headers
    )
    assert empty.json()["total"] == 0


# ----------------------------------------------------------------------
# Caja
# ----------------------------------------------------------------------
def test_cash_registers_report_open_closed_and_difference(
    client, admin_user, admin_headers
):
    opened = client.post(
        "/api/cash/registers/open",
        json={"opening_amount": "100.00"},
        headers=admin_headers,
    )
    assert opened.status_code == 201
    register_id = opened.json()["id"]
    closed = client.post(
        f"/api/cash/registers/{register_id}/close",
        json={"counted_amount": "90.00"},
        headers=admin_headers,
    )
    assert closed.status_code == 200
    difference = closed.json()["difference"]

    reopened = client.post(
        "/api/cash/registers/open",
        json={"opening_amount": "50.00"},
        headers=admin_headers,
    )
    assert reopened.status_code == 201

    data = _get(client, f"{BASE}/cash-registers", admin_headers).json()
    assert data["total"] == 2
    statuses = {r["status_label"] for r in data["items"]}
    assert statuses == {"Cerrada", "Abierta"}
    assert data["totals"]["registers_count"] == 2
    assert data["totals"]["closed_count"] == 1
    assert data["totals"]["difference_sum"] == difference
    assert data["totals"]["opening_sum"] == "150.00"
    open_row = next(r for r in data["items"] if r["status_label"] == "Abierta")
    assert open_row["expected_amount"] is None


def test_cash_registers_report_filter_cashier(
    client, admin_user, cashier_user, admin_headers, cashier_headers
):
    client.post(
        "/api/cash/registers/open",
        json={"opening_amount": "100.00"},
        headers=admin_headers,
    )
    cashier_opened = client.post(
        "/api/cash/registers/open",
        json={"opening_amount": "70.00"},
        headers=cashier_headers,
    )
    assert cashier_opened.status_code == 201

    data = client.get(
        f"{BASE}/cash-registers",
        params={"user_id": str(cashier_user.id)},
        headers=admin_headers,
    ).json()
    assert data["total"] == 1
    assert data["items"][0]["cashier"] == cashier_user.username
    assert data["totals"]["opening_sum"] == "70.00"


# ----------------------------------------------------------------------
# Clientes con crédito
# ----------------------------------------------------------------------
def test_customers_credit_report_balances_and_excludes_non_credit(
    client, admin_user, admin_headers, make_customer
):
    customer, receivable = _credit_customer_with_receivable(
        client, admin_user, make_customer, amount="300.00"
    )
    make_customer(name="Sin crédito", credit_enabled=False)

    paid = client.post(
        f"/api/receivables/{receivable['id']}/payments",
        json={"amount": "100.00", "payment_method": "CASH"},
        headers=admin_headers,
    )
    assert paid.status_code == 201

    data = _get(client, f"{BASE}/customers-credit", admin_headers).json()
    assert data["total"] == 1
    row = data["items"][0]
    assert row["customer_name"] == customer.name
    # El abono de 100 reduce tanto el saldo del cliente como la cuenta.
    assert row["current_balance"] == "200.00"
    assert row["receivable_balance"] == "200.00"
    assert row["open_accounts"] == 1
    totals = data["totals"]
    assert totals["customers_count"] == 1
    assert totals["balance_sum"] == "200.00"
    assert totals["receivable_balance_sum"] == "200.00"


def test_customers_credit_search_q(
    client, admin_user, admin_headers, make_customer
):
    match = make_customer(
        name="Comercial Zully", credit_enabled=True, credit_limit="500.00"
    )
    make_customer(name="Otro Nombre", credit_enabled=True, credit_limit="500.00")

    data = client.get(
        f"{BASE}/customers-credit", params={"q": "zully"}, headers=admin_headers
    ).json()
    assert data["total"] == 1
    assert data["items"][0]["customer_name"] == match.name


# ----------------------------------------------------------------------
# Cuentas por cobrar
# ----------------------------------------------------------------------
def test_receivables_report_statuses_and_totals(
    client, admin_user, admin_headers, make_customer
):
    customer, receivable = _credit_customer_with_receivable(
        client, admin_user, make_customer, amount="400.00"
    )
    _, second = _credit_customer_with_receivable(
        client, admin_user, make_customer, amount="150.00"
    )
    client.post(
        f"/api/receivables/{second['id']}/payments",
        json={"amount": "150.00", "payment_method": "CARD"},
        headers=admin_headers,
    )

    data = _get(client, f"{BASE}/receivables", admin_headers).json()
    assert data["total"] == 2
    assert data["totals"]["original_sum"] == "550.00"
    assert data["totals"]["paid_sum"] == "150.00"
    assert data["totals"]["balance_sum"] == "400.00"

    open_only = client.get(
        f"{BASE}/receivables",
        params={"status": "OPEN"},
        headers=admin_headers,
    ).json()
    assert open_only["total"] == 1
    assert open_only["items"][0]["status_label"] == "Abierta"
    assert open_only["items"][0]["customer_name"] == customer.name
    assert open_only["totals"]["balance_sum"] == "400.00"

    by_customer = client.get(
        f"{BASE}/receivables",
        params={"customer_id": str(customer.id)},
        headers=admin_headers,
    ).json()
    assert by_customer["total"] == 1


# ----------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------
def test_utilities_report_daily_and_monthly_buckets(
    client, admin_user, admin_headers, make_product
):
    product = make_product(purchase_price="50.00", sale_price="100.00")
    _sale(client, admin_user, product, quantity=2)
    _sale(client, admin_user, product, quantity=1)

    daily = client.get(
        f"{BASE}/utilities", params={"period": "DAY"}, headers=admin_headers
    ).json()
    assert daily["report_type"] == "utilities"
    assert daily["total"] >= 1
    # Las ventas se confirman con fecha UTC; el bucket usa esa fecha.
    today_key = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    day_row = next(r for r in daily["items"] if r["period_label"] == today_key)
    assert day_row["sales_count"] == 2
    assert day_row["revenue"] == "300.00"
    assert day_row["cost"] == "150.00"
    assert day_row["profit"] == "150.00"
    assert daily["totals"]["revenue_total"] == "300.00"
    assert daily["totals"]["profit_total"] == "150.00"
    assert daily["totals"]["sales_count_total"] == 2

    monthly = client.get(
        f"{BASE}/utilities", params={"period": "MONTH"}, headers=admin_headers
    ).json()
    month_key = datetime.now(timezone.utc).strftime("%Y-%m")
    month_row = next(
        r for r in monthly["items"] if r["period_label"] == month_key
    )
    assert month_row["revenue"] == "300.00"


def test_utilities_empty_report_zeroed(client, admin_headers):
    data = _get(client, f"{BASE}/utilities", admin_headers).json()
    assert data["items"] == []
    assert data["totals"]["profit_total"] == "0.00"


# ----------------------------------------------------------------------
# Exportación PDF / Excel
# ----------------------------------------------------------------------
@pytest.mark.parametrize("report_key", REPORT_KEYS)
def test_export_pdf_all_reports(client, admin_headers, report_key):
    response = client.get(
        f"{BASE}/{report_key}/export", params={"format": "pdf"}, headers=admin_headers
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=\"reporte-")
    assert re.search(
        rf"reporte-{re.escape(report_key)}-\d{{8}}-\d{{4}}\.pdf", disposition
    )


@pytest.mark.parametrize("report_key", REPORT_KEYS)
def test_export_excel_all_reports(client, admin_headers, report_key):
    response = client.get(
        f"{BASE}/{report_key}/export",
        params={"format": "excel"},
        headers=admin_headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"
    assert re.search(
        rf"reporte-{re.escape(report_key)}-\d{{8}}-\d{{4}}\.xlsx",
        response.headers["content-disposition"],
    )


def test_export_excel_sales_contains_data_and_totals(
    client, admin_user, admin_headers, make_product
):
    product = make_product(sale_price="100.00", stock=30)
    sale = _sale(client, admin_user, product, quantity=2)

    payload = client.get(
        f"{BASE}/sales/export", params={"format": "excel"}, headers=admin_headers
    ).content
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook.active
    cells = [
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    joined = "|".join(cells)
    assert "Reporte de ventas" in joined
    assert sale["sale_number"] in joined
    assert "TOTALES" in cells
    total_dec = Decimal(sale["total"])
    numeric_values = [
        Decimal(str(cell.value))
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float))
    ]
    assert total_dec in numeric_values  # fila de datos y totales contienen el monto


def test_export_excel_respects_filters_applied_header(
    client, admin_user, admin_headers, make_product
):
    product = make_product(stock=30)
    _sale(client, admin_user, product, quantity=2, payment_method="CARD")
    payload = client.get(
        f"{BASE}/sales/export",
        params={"format": "excel", "payment_method": "CARD"},
        headers=admin_headers,
    ).content
    workbook = load_workbook(io.BytesIO(payload))
    header_cell = workbook.active.cell(row=3, column=1).value
    assert header_cell and "Método de pago" in header_cell


def test_export_pdf_empty_report_still_valid(client, admin_headers):
    response = client.get(
        f"{BASE}/sales/export", params={"format": "pdf"}, headers=admin_headers
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_export_ignores_pagination_and_exports_all_rows(
    client, admin_user, admin_headers, make_product
):
    product = make_product(stock=100, minimum_stock=0)
    for _ in range(3):
        _sale(client, admin_user, product, quantity=1)

    limited = client.get(
        f"{BASE}/sales", params={"limit": 2}, headers=admin_headers
    ).json()
    assert len(limited["items"]) == 2

    payload = client.get(
        f"{BASE}/sales/export", params={"format": "excel"}, headers=admin_headers
    ).content
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook.active
    numbers_in_sheet = {
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value and str(cell.value).startswith("VTA-")
    }
    assert len(numbers_in_sheet) == 3


def test_format_money_includes_currency_symbol():
    assert format_money("1234.5", symbol="$", decimals=2) == "$ 1,234.50"
    assert format_money("1234.5") == "1,234.50"
    assert format_money("10", symbol="$", decimals=0) == "$ 10"
    assert format_money(None, symbol="$") == "-"


def test_currency_number_format_builds_excel_pattern():
    assert currency_number_format("$", 2) == '"$" #,##0.00'
    assert currency_number_format(None, 0) == "#,##0"
    assert currency_number_format("COP", 0) == '"COP" #,##0'


def test_export_excel_applies_configured_currency_symbol(
    client, admin_user, admin_headers, make_product
):
    make_product(sale_price="100.00", stock=10)
    client.put(
        "/api/settings",
        json={"currency_code": "USD", "currency_symbol": "US$"},
        headers=admin_headers,
    )
    payload = client.get(
        f"{BASE}/sales/export", params={"format": "excel"}, headers=admin_headers
    ).content
    workbook = load_workbook(io.BytesIO(payload))
    money_formats = {
        cell.number_format
        for row in workbook.active.iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float))
    }
    assert any("US$" in fmt for fmt in money_formats)
