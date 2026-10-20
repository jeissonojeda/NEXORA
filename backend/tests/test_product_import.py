"""Tests de importación masiva de productos (Fase 2/4)."""

import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models.category import Category
from app.models.inventory import InventoryMovement
from app.models.product import Product
from app.models.product_import import ProductImportBatch
from app.models.supplier import Supplier
from app.services.product_import_service import COLUMNS
from tests.conftest import auth_headers

HEADER = [c.upper() for c in COLUMNS]


def _xlsx(rows: list[list], header: list | None = None) -> bytes:
    """Construye un XLSX en memoria con hoja 'Productos'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(header if header is not None else HEADER)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row(
    nombre="Producto Import",
    categoria="Abarrotes",
    codigo="",
    precio_compra="",
    precio_venta="",
    stock_inicial="",
    stock_minimo="",
    activo="",
    **override,
) -> list:
    """Fila en el orden canónico de COLUMNS."""
    values = {
        "codigo_barras": codigo,
        "nombre": nombre,
        "categoria": categoria,
        "descripcion": "",
        "proveedor": "",
        "unidad": "",
        "precio_compra": precio_compra,
        "margen_utilidad_pct": "",
        "precio_venta": precio_venta,
        "impuesto": "",
        "stock_inicial": stock_inicial,
        "stock_minimo": stock_minimo,
        "activo": activo,
    }
    values.update(override)
    return [values[c] for c in COLUMNS]


def _preview(client, user, content, name="import.xlsx", **form):
    return client.post(
        "/api/products/import/preview",
        files={"file": (name, io.BytesIO(content))},
        headers=auth_headers(user),
    )


def _confirm(
    client,
    user,
    content,
    policy="SKIP_NEW",
    name="import.xlsx",
    overrides: list[dict] | str | None = None,
):
    data = {"duplicate_policy": policy}
    if overrides is not None:
        data["overrides"] = (
            overrides if isinstance(overrides, str) else json.dumps(overrides)
        )
    return client.post(
        "/api/products/import/confirm",
        files={"file": (name, io.BytesIO(content))},
        data=data,
        headers=auth_headers(user),
    )


# ----------------------------------------------------------------------
# Plantilla y plantillas de permiso
# ----------------------------------------------------------------------
def test_template_download_xlsx(client, admin_user):
    resp = client.get(
        "/api/products/import/template", headers=auth_headers(admin_user)
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Productos" in wb.sheetnames
    sheet = wb["Productos"]
    assert sheet.cell(row=1, column=1).value == COLUMNS[0].upper()


def test_import_requires_permission(client, cashier_user):
    resp = client.get(
        "/api/products/import/template", headers=auth_headers(cashier_user)
    )
    assert resp.status_code == 403


def test_inventory_manager_can_access(client, inventory_manager_user):
    resp = client.get(
        "/api/products/import/template",
        headers=auth_headers(inventory_manager_user),
    )
    assert resp.status_code == 200


# ----------------------------------------------------------------------
# Preview
# ----------------------------------------------------------------------
def test_preview_valid_row(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx([_row(nombre="Arroz", categoria="Abarrotes", precio_venta="5000")])
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_rows"] == 1
    assert data["valid"] == 1
    assert data["errors"] == 0
    assert data["rows"][0]["status"] == "VALID"


def test_preview_missing_required_columns(client, admin_user):
    content = _xlsx([["Solo", "Dos"]], header=["nombre", "categoria", "extra"])
    # header sin 'categoria' real -> el servicio valida columnas obligatorias
    content = _xlsx([["x"]], header=["nombre"])
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_MISSING_COLUMNS"


def test_preview_unknown_category_rejected(client, admin_user):
    content = _xlsx([_row(nombre="X", categoria="NoExiste")])
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    row = resp.json()["rows"][0]
    assert row["status"] == "ERROR"
    assert any("categoria" in m for m in row["messages"])


def test_preview_invalid_price_rejected(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx([_row(nombre="X", categoria="Abarrotes", precio_venta="abc")])
    resp = _preview(client, admin_user, content)
    row = resp.json()["rows"][0]
    assert row["status"] == "ERROR"
    assert any("precio_venta" in m for m in row["messages"])


def test_preview_duplicate_barcode_in_file(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(nombre="A", categoria="Abarrotes", codigo="777", precio_venta="100"),
            _row(nombre="B", categoria="Abarrotes", codigo="777", precio_venta="100"),
        ]
    )
    resp = _preview(client, admin_user, content)
    rows = resp.json()["rows"]
    assert rows[0]["status"] == "VALID"
    assert rows[1]["status"] == "ERROR"
    assert any("duplicado" in m for m in rows[1]["messages"])


def test_preview_unsupported_format(client, admin_user):
    resp = _preview(client, admin_user, b"data", name="productos.txt")
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_UNSUPPORTED_FORMAT"


def test_preview_empty_file(client, admin_user):
    resp = client.post(
        "/api/products/import/preview",
        files={"file": ("vacio.xlsx", io.BytesIO(b""))},
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_EMPTY_FILE"


# ----------------------------------------------------------------------
# Confirm: creación, stock e idempotencia de stock
# ----------------------------------------------------------------------
def test_confirm_creates_product_with_initial_stock(
    client, admin_user, db_session, make_category
):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(
                nombre="Nuevo Producto",
                categoria="Abarrotes",
                codigo="9990001",
                precio_compra="2000",
                precio_venta="3000",
                stock_inicial="50",
                stock_minimo="5",
            )
        ]
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["created"] == 1
    assert data["rejected"] == 0

    product = db_session.scalars(
        select(Product).where(Product.barcode == "9990001")
    ).first()
    assert product is not None
    assert product.stock == 50
    assert product.minimum_stock == 5

    movements = db_session.scalars(
        select(InventoryMovement).where(InventoryMovement.product_id == product.id)
    ).all()
    assert len(movements) == 1
    assert movements[0].quantity == 50
    assert movements[0].reference_type == "IMPORT"


def test_confirm_skips_duplicate_with_skip_new(
    client, admin_user, db_session, make_product, make_category
):
    make_category(name="Abarrotes")
    existing = make_product(
        name="Existente", barcode="5550001", stock=10, **{"category": None}
    )
    content = _xlsx(
        [
            _row(
                nombre="Intento Dup",
                categoria="Abarrotes",
                codigo="5550001",
                stock_inicial="999",
            )
        ]
    )
    resp = _confirm(client, admin_user, content, policy="SKIP_NEW")
    data = resp.json()
    assert data["skipped"] == 1
    assert data["created"] == 0

    db_session.refresh(existing)
    # El stock del producto existente NUNCA se sobreescribe.
    assert existing.stock == 10
    assert existing.name == "Existente"


def test_confirm_updates_duplicate_with_update_policy(
    client, admin_user, db_session, make_product, make_category
):
    make_category(name="Abarrotes")
    existing = make_product(
        name="Existente", barcode="5550002", stock=10, sale_price="100.00"
    )
    content = _xlsx(
        [
            _row(
                nombre="Nombre Actualizado",
                categoria="Abarrotes",
                codigo="5550002",
                precio_venta="777",
                stock_inicial="999",
            )
        ]
    )
    resp = _confirm(client, admin_user, content, policy="UPDATE_EXISTING")
    data = resp.json()
    assert data["updated"] == 1

    db_session.refresh(existing)
    assert existing.name == "Existente"  # el nombre no se sobrescribe
    assert existing.sale_price == 777
    assert existing.stock == 10  # stock intacto


def test_confirm_rejects_invalid_policy(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx([_row(nombre="X", categoria="Abarrotes")])
    resp = _confirm(client, admin_user, content, policy="INVENTADA")
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "INVALID_DUPLICATE_POLICY"


def test_confirm_registers_batch_audit(
    client, admin_user, db_session, make_category
):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(nombre="A", categoria="Abarrotes", codigo="B1"),
            _row(nombre="B", categoria="NoExiste"),  # rechazada
        ]
    )
    resp = _confirm(client, admin_user, content)
    data = resp.json()
    assert data["created"] == 1
    assert data["rejected"] == 1

    batch = db_session.get(ProductImportBatch, __import__("uuid").UUID(data["batch_id"]))
    assert batch is not None
    assert batch.created_count == 1
    assert batch.rejected_count == 1
    assert batch.created_by == admin_user.id


# ----------------------------------------------------------------------
# Historial e informes
# ----------------------------------------------------------------------
def test_list_batches(client, admin_user, db_session, make_category):
    make_category(name="Abarrotes")
    _confirm(client, admin_user, _xlsx([_row(nombre="A", categoria="Abarrotes")]))
    resp = client.get(
        "/api/products/import/batches", headers=auth_headers(admin_user)
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] >= 1


def test_error_report_download(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(nombre="Bueno", categoria="Abarrotes"),
            _row(nombre="Malo", categoria="NoExiste"),
        ]
    )
    batch_id = _confirm(client, admin_user, content).json()["batch_id"]
    resp = client.get(
        f"/api/products/import/batches/{batch_id}/errors",
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.title == "Rechazadas"


def test_error_report_no_errors(client, admin_user, make_category):
    make_category(name="Abarrotes")
    batch_id = _confirm(
        client, admin_user, _xlsx([_row(nombre="A", categoria="Abarrotes")])
    ).json()["batch_id"]
    resp = client.get(
        f"/api/products/import/batches/{batch_id}/errors",
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 404
    assert resp.json()["error_code"] == "IMPORT_BATCH_NO_ERRORS"


# ----------------------------------------------------------------------
# CSV y sanitización
# ----------------------------------------------------------------------
def test_confirm_csv_import(client, admin_user, db_session, make_category):
    make_category(name="Abarrotes")
    csv_text = (
        "nombre,categoria,precio_venta,stock_inicial\n"
        "CsvProduct,Abarrotes,1500,7\n"
    )
    resp = _confirm(
        client, admin_user, csv_text.encode("utf-8"), name="productos.csv"
    )
    assert resp.status_code == 200
    assert resp.json()["created"] == 1
    product = db_session.scalars(
        select(Product).where(Product.name == "CsvProduct")
    ).first()
    assert product is not None
    assert product.stock == 7


def test_csv_formula_injection_sanitized(client, admin_user, db_session, make_category):
    """El informe de errores exportado neutraliza fórmulas (=, +, -, @)."""
    make_category(name="Abarrotes")
    # En CSV el texto '=' se conserva literal (en XLSX sería una fórmula).
    csv_text = "nombre,categoria\n=CMD(),NoExiste\n"
    batch_id = _confirm(
        client, admin_user, csv_text.encode("utf-8"), name="p.csv"
    ).json()["batch_id"]
    resp = client.get(
        f"/api/products/import/batches/{batch_id}/errors",
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    # Columna 'nombre' en posición 4 (1: fila, 2: errores, 3: codigo_barras, 4: nombre).
    exported = str(ws.cell(row=2, column=4).value)
    assert not exported.startswith("=")
    assert exported.startswith("'")


# ----------------------------------------------------------------------
# Plantilla sin datos de ejemplo y sugerencias de referencias
# ----------------------------------------------------------------------
def test_template_has_only_headers(client, admin_user):
    """La plantilla no incluye filas de ejemplo con entidades inexistentes."""
    resp = client.get(
        "/api/products/import/template", headers=auth_headers(admin_user)
    )
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Instrucciones" in wb.sheetnames
    sheet = wb["Productos"]
    assert sheet.max_row == 1

    values = {
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Abarrotes" not in values
    assert "Distribuidora X" not in values
    assert "IVA 19%" not in values


def test_preview_suggests_accent_insensitive_category(
    client, admin_user, make_category
):
    make_category(name="Lácteos")
    content = _xlsx([_row(nombre="Leche", categoria="Lacteos")])
    resp = _preview(client, admin_user, content)
    row = resp.json()["rows"][0]
    assert row["status"] == "ERROR"
    assert row["data"]["category_id"] is None
    assert row["data"]["suggested_category"] == "Lácteos"
    assert any("quiso decir" in m for m in row["messages"])


def test_preview_override_resolves_reference(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx(
        [_row(nombre="Arroz", categoria="NoExiste", precio_venta="100")]
    )
    resp = client.post(
        "/api/products/import/preview",
        files={"file": ("import.xlsx", io.BytesIO(content))},
        data={
            "overrides": json.dumps(
                [{"row_number": 2, "category": "Abarrotes"}]
            )
        },
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 200
    row = resp.json()["rows"][0]
    assert row["status"] == "VALID"
    assert row["data"]["category_name"] == "Abarrotes"
    assert row["data"]["category_id"] is not None


def test_confirm_with_override_creates_corrected_product(
    client, admin_user, db_session, make_category
):
    make_category(name="Abarrotes")
    content = _xlsx(
        [_row(nombre="Arroz", categoria="NoExiste", precio_venta="100")]
    )
    resp = _confirm(
        client,
        admin_user,
        content,
        overrides=[{"row_number": 2, "category": "Abarrotes"}],
    )
    assert resp.status_code == 200
    assert resp.json()["created"] == 1

    product = db_session.scalars(
        select(Product).where(Product.name == "Arroz")
    ).first()
    assert product is not None
    assert product.category.name == "Abarrotes"


def test_confirm_invalid_overrides_rejected(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx([_row(nombre="X", categoria="Abarrotes")])
    resp = _confirm(client, admin_user, content, overrides="{no-es-json")
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_INVALID_OVERRIDES"


def test_import_entities_lists_catalog(client, admin_user, make_category):
    make_category(name="Abarrotes")
    resp = client.get(
        "/api/products/import/entities", headers=auth_headers(admin_user)
    )
    assert resp.status_code == 200
    data = resp.json()
    assert any(c["name"] == "Abarrotes" for c in data["categories"])
    assert set(data.keys()) == {"categories", "suppliers", "taxes"}


def test_import_entities_requires_permission(client, cashier_user):
    resp = client.get(
        "/api/products/import/entities", headers=auth_headers(cashier_user)
    )
    assert resp.status_code == 403


# ----------------------------------------------------------------------
# Duplicados sin código de barras (reintentos seguros)
# ----------------------------------------------------------------------
def test_preview_duplicate_name_without_barcode_is_warning(
    client, admin_user, make_category, make_product
):
    category = make_category(name="Abarrotes")
    make_product(name="Arroz", barcode=None, category=category)
    content = _xlsx(
        [_row(nombre="Arroz", categoria="Abarrotes", precio_venta="100")]
    )
    row = _preview(client, admin_user, content).json()["rows"][0]
    assert row["status"] == "WARNING"
    assert row["existing_product_id"] is not None
    assert any("ya registrado" in m for m in row["messages"])


def test_confirm_skips_name_duplicate_without_barcode(
    client, admin_user, db_session, make_category, make_product
):
    category = make_category(name="Abarrotes")
    make_product(name="Arroz", barcode=None, category=category)
    content = _xlsx(
        [_row(nombre="Arroz", categoria="Abarrotes", precio_venta="100")]
    )
    resp = _confirm(client, admin_user, content, policy="SKIP_NEW")
    data = resp.json()
    assert data["skipped"] == 1
    assert data["created"] == 0

    matches = db_session.scalars(
        select(Product).where(Product.name == "Arroz")
    ).all()
    assert len(matches) == 1


def test_preview_duplicate_name_in_file_rejected(
    client, admin_user, make_category
):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(nombre="Arroz", categoria="Abarrotes", precio_venta="100"),
            _row(nombre="Arroz", categoria="Abarrotes", precio_venta="100"),
        ]
    )
    rows = _preview(client, admin_user, content).json()["rows"]
    assert rows[0]["status"] == "VALID"
    assert rows[1]["status"] == "ERROR"
    assert any("duplicado dentro del archivo" in m for m in rows[1]["messages"])


# ----------------------------------------------------------------------
# Plantilla multi-hoja y creación de entidades desde el archivo
# ----------------------------------------------------------------------
def _multi_xlsx(
    product_rows: list[list],
    suppliers: list[list] | None = None,
    categories: list[list] | None = None,
    taxes: list[list] | None = None,
    product_header: list | None = None,
) -> bytes:
    """Construye un XLSX con las hojas Productos/Proveedores/Categorías."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(product_header if product_header is not None else HEADER)
    for row in product_rows:
        ws.append(row)
    if suppliers is not None:
        ws_sup = wb.create_sheet("Proveedores")
        ws_sup.append(["NIT", "NOMBRE", "TELEFONO", "EMAIL", "DIRECCION", "CONTACTO", "NOTAS"])
        for row in suppliers:
            ws_sup.append(row)
    if categories is not None:
        ws_cat = wb.create_sheet("Categorias")
        ws_cat.append(["NOMBRE", "DESCRIPCION", "ACTIVA"])
        for row in categories:
            ws_cat.append(row)
    if taxes is not None:
        ws_tax = wb.create_sheet("Impuestos")
        ws_tax.append(["NOMBRE", "TASA", "ACTIVO"])
        for row in taxes:
            ws_tax.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_template_has_all_sheets(client, admin_user):
    resp = client.get(
        "/api/products/import/template", headers=auth_headers(admin_user)
    )
    wb = load_workbook(io.BytesIO(resp.content))
    # Sheet lookup is accent/case-insensitive in the importer; verify the
    # canonical set is present (uppercase for the data sheets).
    names = {n.upper() for n in wb.sheetnames}
    for expected in ["PRODUCTOS", "PROVEEDORES", "CATEGORIAS", "IMPUESTOS", "INSTRUCCIONES"]:
        assert expected in names


def test_preview_creates_declared_category(client, admin_user):
    content = _multi_xlsx(
        product_rows=[
            _row(
                nombre="Arroz",
                categoria="Abarrotes",
                precio_venta="100",
                **{"proveedor": ""},
            )
        ],
        categories=[["Abarrotes", "Supermercado", "SI"]],
    )
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["entities"]["categories"]["new"] == 1
    row = data["rows"][0]
    assert row["status"] in ("VALID", "WARNING")
    assert row["data"]["category_id"] is None
    assert row["data"]["category_declared"] is True


def test_confirm_creates_declared_category_and_product(
    client, admin_user, db_session
):
    content = _multi_xlsx(
        product_rows=[_row(nombre="Arroz", categoria="Abarrotes", precio_venta="100")],
        categories=[["Abarrotes", "Supermercado", "SI"]],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["created"] == 1
    assert data["categories_created"] == 1

    category = db_session.scalars(
        select(Category).where(Category.name == "Abarrotes")
    ).first()
    assert category is not None
    product = db_session.scalars(
        select(Product).where(Product.name == "Arroz")
    ).first()
    assert product is not None
    assert product.category_id == category.id


def test_confirm_creates_declared_supplier(client, admin_user, db_session, make_category):
    make_category(name="Abarrotes")
    content = _multi_xlsx(
        product_rows=[
            _row(
                nombre="Arroz",
                categoria="Abarrotes",
                precio_venta="100",
                proveedor="Distribuidora X",
            )
        ],
        suppliers=[
            ["900123456-1", "Distribuidora X", "3001234567", "v@x.com", "Calle 1", "Carlos", ""]
        ],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    assert resp.json()["suppliers_created"] == 1

    supplier = db_session.scalars(
        select(Supplier).where(Supplier.name == "Distribuidora X")
    ).first()
    assert supplier is not None
    product = db_session.scalars(select(Product).where(Product.name == "Arroz")).first()
    assert product.supplier_id == supplier.id


def test_confirm_reuses_existing_supplier(client, admin_user, db_session, make_category, make_supplier):
    make_category(name="Abarrotes")
    existing = make_supplier(name="Distribuidora X")
    content = _multi_xlsx(
        product_rows=[
            _row(
                nombre="Arroz",
                categoria="Abarrotes",
                precio_venta="100",
                proveedor="Distribuidora X",
            )
        ],
        suppliers=[["", "Distribuidora X", "", "", "", "", ""]],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    assert resp.json()["suppliers_created"] == 0

    suppliers = db_session.scalars(
        select(Supplier).where(Supplier.name == "Distribuidora X")
    ).all()
    assert len(suppliers) == 1
    assert suppliers[0].id == existing.id


def test_confirm_new_tax_rejected(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _multi_xlsx(
        product_rows=[
            _row(nombre="Arroz", categoria="Abarrotes", precio_venta="100", impuesto="IVA 99%")
        ],
        taxes=[["IVA 99%", "99", "SI"]],
    )
    resp = _preview(client, admin_user, content)
    row = resp.json()["rows"][0]
    assert row["status"] == "ERROR"
    assert any("impuesto" in m for m in row["messages"])
    assert resp.json()["entities"]["taxes"]["new"] == 1


def test_confirm_existing_tax_used(client, admin_user, db_session, make_category):
    from app.models.tax import Tax

    make_category(name="Abarrotes")
    tax = Tax(name="IVA 19%", rate="19", is_active=True)
    db_session.add(tax)
    db_session.commit()
    db_session.refresh(tax)
    content = _multi_xlsx(
        product_rows=[
            _row(nombre="Arroz", categoria="Abarrotes", precio_venta="100", impuesto="IVA 19%")
        ],
        taxes=[["IVA 19%", "19", "SI"]],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    product = db_session.scalars(select(Product).where(Product.name == "Arroz")).first()
    assert product.tax_id == tax.id


def test_inventory_manager_can_create_entities(
    client, inventory_manager_user, db_session, make_category
):
    """El encargado de inventario tiene suppliers.create/categories.create."""
    make_category(name="Abarrotes")
    content = _multi_xlsx(
        product_rows=[
            _row(
                nombre="Arroz",
                categoria="Abarrotes",
                precio_venta="100",
                proveedor="Distribuidora X",
            )
        ],
        suppliers=[[None, "Distribuidora X", "", "", "", "", ""]],
    )
    resp = _confirm(client, inventory_manager_user, content)
    assert resp.status_code == 200
    assert resp.json()["suppliers_created"] == 1


def test_confirm_entity_permission_denied(
    client, admin_user, db_session, make_category, monkeypatch
):
    """Sin permiso de crear categorías, la importación falla sin escribir."""
    make_category(name="Abarrotes")
    content = _multi_xlsx(
        product_rows=[_row(nombre="Arroz", categoria="Abarrotes", precio_venta="100")],
        categories=[["Abarrotes2", "", "SI"]],
    )

    from app.services.rbac_service import RbacService

    original = RbacService.user_has_permission

    def deny(self, user, code):
        if code in ("categories.create", "suppliers.create"):
            return False
        return original(self, user, code)

    monkeypatch.setattr(RbacService, "user_has_permission", deny)

    # Forzar la referencia a una categoría declarada (nueva).
    content = _multi_xlsx(
        product_rows=[
            _row(nombre="Arroz", categoria="Abarrotes2", precio_venta="100")
        ],
        categories=[["Abarrotes2", "", "SI"]],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 403
    assert resp.json()["error_code"] == "IMPORT_ENTITY_PERMISSION_DENIED"
    created = db_session.scalars(
        select(Category).where(Category.name == "Abarrotes2")
    ).all()
    assert created == []


def test_confirm_rollback_on_product_error(client, admin_user, db_session):
    """Si un producto falla al crear, no quedan categorías parciales."""
    from app.core.exceptions import AppError
    from unittest.mock import patch

    content = _multi_xlsx(
        product_rows=[_row(nombre="Arroz", categoria="NuevaCat", precio_venta="100")],
        categories=[["NuevaCat", "", "SI"]],
    )
    with patch(
        "app.services.product_import_service.ProductImportService._create_one",
        side_effect=AppError(status_code=422, error_code="X", message="boom"),
    ):
        resp = _confirm(client, admin_user, content)
    assert resp.status_code == 422
    # La categoría declarada NO debe haberse creado (rollback total).
    created = db_session.scalars(
        select(Category).where(Category.name == "NuevaCat")
    ).all()
    assert created == []


# ----------------------------------------------------------------------
# Robustez: archivos inválidos, límites y formato
# ----------------------------------------------------------------------
def test_preview_empty_file_rejected(client, admin_user):
    resp = client.post(
        "/api/products/import/preview",
        files={"file": ("import.xlsx", io.BytesIO(b""))},
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_EMPTY_FILE"


def test_preview_missing_columns_rejected(client, admin_user):
    # Falta 'categoria' (columna obligatoria).
    content = _xlsx(
        [["Arroz", "100"]], header=["NOMBRE", "PRECIO_VENTA"]
    )
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_MISSING_COLUMNS"
    assert "CATEGORIA" in resp.json()["detail"]


def test_preview_unsupported_extension_rejected(client, admin_user):
    resp = client.post(
        "/api/products/import/preview",
        files={"file": ("import.txt", io.BytesIO(b"nombre,categoria"))},
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_UNSUPPORTED_FORMAT"


def test_preview_invalid_xlsx_content_rejected(client, admin_user):
    resp = _preview(client, admin_user, b"no es un xlsx de verdad")
    assert resp.status_code == 422
    assert resp.json()["error_code"] == "IMPORT_INVALID_FILE"


def test_preview_too_many_rows_rejected(client, admin_user, make_category):
    make_category(name="Abarrotes")
    from app.services.product_import_service import MAX_ROWS

    rows = [
        _row(nombre=f"Producto {i}", categoria="Abarrotes", precio_venta="100")
        for i in range(MAX_ROWS + 1)
    ]
    content = _xlsx(rows)
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 413
    assert resp.json()["error_code"] == "IMPORT_TOO_MANY_ROWS"


def test_preview_file_too_large_rejected(client, admin_user):
    from app.services.product_import_service import MAX_FILE_BYTES

    big = io.BytesIO()
    big.write(b"x" * (MAX_FILE_BYTES + 10))
    resp = client.post(
        "/api/products/import/preview",
        files={"file": ("import.xlsx", io.BytesIO(big.getvalue()))},
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 413
    assert resp.json()["error_code"] == "IMPORT_FILE_TOO_LARGE"


def test_preview_format_errors_reported_in_rows(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx(
        [
            _row(
                nombre="Arroz",
                categoria="Abarrotes",
                precio_venta="no-es-numero",
                stock_inicial="abc",
            )
        ]
    )
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    row = resp.json()["rows"][0]
    assert row["status"] == "ERROR"
    assert any("precio_venta" in m for m in row["messages"])
    assert any("stock_inicial" in m for m in row["messages"])


def test_preview_special_characters_preserved(client, admin_user, make_category):
    make_category(name="Lácteos")
    content = _xlsx(
        [_row(nombre="Queso Münster Íñigo", categoria="Lácteos", precio_venta="100")]
    )
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    row = resp.json()["rows"][0]
    assert row["data"]["name"] == "Queso Münster Íñigo"


def test_preview_formula_injection_neutralized(client, admin_user, make_category):
    make_category(name="Abarrotes")
    content = _xlsx(
        [_row(nombre="=SUM(1+1)", categoria="Abarrotes", precio_venta="100")]
    )
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    # La cadena no debe conservar el prefijo de fórmula peligroso.
    name = resp.json()["rows"][0]["data"]["name"]
    assert not name.startswith("=")


def test_import_endpoint_not_found_returns_404(client, admin_user):
    """Un endpoint inexistente responde 404 con error_code NOT_FOUND."""
    resp = client.post(
        "/api/products/import/no-existe",
        headers=auth_headers(admin_user),
    )
    assert resp.status_code == 404
    assert resp.json()["error_code"] == "NOT_FOUND"


# ----------------------------------------------------------------------
# Escenario real: plantilla de 6 productos con catálogos que NO existen
# ----------------------------------------------------------------------
def test_realistic_template_with_missing_catalogs_returns_errors(
    client, admin_user
):
    """Un archivo que declara categoría/proveedor/impuesto inexistentes no
    debe fallar con 500; devuelve filas en ERROR con mensajes claros."""
    product_rows = [
        _row(nombre="Arroz 1kg", categoria="Abarrotes", proveedor="Distribuidora X", impuesto="IVA 19%", precio_venta="100"),
        _row(nombre="Aceite vegetal 900 ml", categoria="Abarrotes", proveedor="Distribuidora X", impuesto="IVA 19%", precio_venta="100"),
        _row(nombre="Leche entera 1L", categoria="Lácteos", proveedor="Distribuidora Y", impuesto="IVA 19%", precio_venta="100"),
        _row(nombre="Jabón en polvo", categoria="Aseo", proveedor="Distribuidora Z", impuesto="IVA 19%", precio_venta="100"),
        _row(nombre="Gaseosa 1.5L", categoria="Bebidas", proveedor="Distribuidora X", impuesto="IVA 19%", precio_venta="100"),
        _row(nombre="Pan tajado", categoria="Panadería", proveedor="Distribuidora Y", impuesto="IVA 19%", precio_venta="100"),
    ]
    content = _xlsx(product_rows)
    resp = _preview(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_rows"] == 6
    assert data["errors"] == 6
    # Cada fila reporta la categoría inexistente sin romper la respuesta.
    assert all(row["status"] == "ERROR" for row in data["rows"])
    assert all(
        any("categoria" in m or "Categoría" in m for m in row["messages"])
        for row in data["rows"]
    )


def test_realistic_template_creates_declared_entities(
    client, admin_user, db_session
):
    """Con hojas auxiliares completas, se crean categorías/proveedores y se
    valida el impuesto existente; los productos importan correctamente."""
    from app.models.tax import Tax

    tax = Tax(name="IVA 19%", rate="19", is_active=True)
    db_session.add(tax)
    db_session.commit()

    content = _multi_xlsx(
        product_rows=[
            _row(nombre="Arroz 1kg", categoria="Abarrotes", proveedor="Distribuidora X", impuesto="IVA 19%", precio_venta="100"),
            _row(nombre="Aceite 900 ml", categoria="Abarrotes", proveedor="Distribuidora X", impuesto="IVA 19%", precio_venta="100"),
        ],
        suppliers=[[None, "Distribuidora X", "", "", "", "", ""]],
        categories=[["Abarrotes", "Supermercado", "SI"]],
        taxes=[["IVA 19%", "19", "SI"]],
    )
    resp = _confirm(client, admin_user, content)
    assert resp.status_code == 200
    data = resp.json()
    assert data["created"] == 2
    assert data["categories_created"] == 1
    assert data["suppliers_created"] == 1
    assert data["rejected"] == 0

