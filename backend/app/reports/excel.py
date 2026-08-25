"""Generación de documentos Excel (xlsx) para reportes.

Usa openpyxl (tecnología definida en docs/ARCHITECTURE.md §2). Hoja única
con: fila de título, fecha/hora de generación, filtros aplicados,
encabezados con formato, datos tabulares (números reales para montos) y
fila(s) de totales. Renderizador genérico sin lógica de dominio.
"""

import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
_TOTALS_FILL = PatternFill("solid", fgColor="E2E8F0")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

MONEY_FORMAT = "#,##0.00"


def build_excel(
    title: str,
    generated_at: datetime,
    filters_applied: list[tuple[str, str]],
    columns: list[dict],
    rows: list[dict],
    totals_rows: list[tuple[str, dict]],
) -> bytes:
    """Construye el xlsx y retorna los bytes.

    - ``columns``: [{"key": ..., "label": ..., "type": "text"|"money"|
      "int"|"datetime"|"date"}].
    - ``rows``: dicts key→valor (Decimal para money).
    - ``totals_rows``: pares (etiqueta, dict key→valor).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"

    bold = Font(bold=True)
    header_font = Font(bold=True, color="FFFFFF")

    row_index = 1
    title_cell = sheet.cell(row=row_index, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    row_index += 1
    sheet.cell(
        row=row_index,
        column=1,
        value=f"Generado: {generated_at.strftime('%d/%m/%Y %H:%M')}",
    )
    row_index += 1
    if filters_applied:
        applied = " · ".join(
            f"{label}: {value}" if value else label
            for label, value in filters_applied
        )
        sheet.cell(row=row_index, column=1, value=f"Filtros: {applied}")
        row_index += 1
    row_index += 1

    # Encabezados
    header_row = row_index
    for col_offset, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=col_offset, value=column["label"])
        cell.font = header_font
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row_index += 1

    def write_value(cell, value: object, column_type: str) -> None:
        if column_type == "money":
            amount = None if value is None else Decimal(str(value))
            cell.value = float(amount) if amount is not None else None
            cell.number_format = MONEY_FORMAT
        elif column_type == "int":
            cell.value = int(value) if value is not None else None
        elif column_type == "datetime":
            cell.value = value.replace(tzinfo=None) if value is not None else None
            cell.number_format = "DD/MM/YYYY HH:MM"
        elif column_type == "date":
            cell.value = value.replace(tzinfo=None) if value is not None else None
            cell.number_format = "DD/MM/YYYY"
        else:
            cell.value = "" if value is None else str(value)
        cell.border = _BORDER

    first_data_row = row_index
    for data_row in rows:
        for col_offset, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=col_offset)
            write_value(cell, data_row.get(column["key"]), column.get("type", "text"))
        row_index += 1

    # Totales: etiqueta en la primera columna, valores en las restantes.
    for label, totals_map in totals_rows:
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        label_cell.font = bold
        label_cell.fill = _TOTALS_FILL
        label_cell.border = _BORDER
        for col_offset, column in enumerate(columns[1:], start=2):
            cell = sheet.cell(row=row_index, column=col_offset)
            cell.fill = _TOTALS_FILL
            cell.font = bold
            write_value(cell, totals_map.get(column["key"]), column.get("type", "text"))
        row_index += 1

    # Anchos de columna razonables
    widths: dict[int, int] = {}
    for col_offset, column in enumerate(columns, start=1):
        header_width = max(len(str(column["label"])) + 4, 10)
        sample_width = max(
            (
                len(str(row_data.get(column["key"]) or ""))
                for row_data in rows[:50]
            ),
            default=0,
        )
        widths[col_offset] = min(max(header_width, sample_width + 2), 42)
    for col_offset, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_offset)].width = width

    # Filtro automático y paneles congelados en la tabla de datos
    if rows:
        last_col = get_column_letter(len(columns))
        last_row = first_data_row + len(rows) - 1
        sheet.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    sheet.freeze_panes = f"A{first_data_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
